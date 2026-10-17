"""
Builds the report in two formats from the same text (report_text.py):

    report/Speech_Phase_STOI_Report.pdf    (ReportLab, Times New Roman, A4)
    report/Speech_Phase_STOI_Report.docx   (Word; opens directly in Google Docs)

    python make_report.py --name "Sarthak Tripathi" --usn 1MV23EC096

All numbers come from outputs/tables and outputs/results.json, so run run_all.py first.
The text is built twice so that figure and table numbers can be referred to before they appear.
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path
from xml.sax.saxutils import escape

import matplotlib
import numpy as np
import pandas as pd
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (CondPageBreak, Image, KeepTogether, PageBreak, Paragraph, SimpleDocTemplate,
                                Spacer, Table, TableStyle)

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "outputs"
FIG = OUT / "figures"
TAB = OUT / "tables"
REP = ROOT / "report"

WINFONTS = Path("C:/Windows/Fonts")
_dv = Path(matplotlib.__file__).parent / "mpl-data" / "fonts" / "ttf"
for name, f in [("TNR", WINFONTS / "times.ttf"), ("TNR-B", WINFONTS / "timesbd.ttf"),
                ("TNR-I", WINFONTS / "timesi.ttf"), ("TNR-BI", WINFONTS / "timesbi.ttf"),
                ("DVS", _dv / "DejaVuSerif.ttf")]:
    pdfmetrics.registerFont(TTFont(name, str(f)))
pdfmetrics.registerFontFamily("TNR", normal="TNR", bold="TNR-B", italic="TNR-I", boldItalic="TNR-BI")

# Times New Roman has no glyph for these symbols, so the PDF draws them with DejaVu Serif
_MISSING = re.compile("([\u2208\u2209\u2220\u2218])")


def fix(text: str) -> str:
    return _MISSING.sub(r'<font name="DVS">\1</font>', text)


GREY = colors.HexColor("#444444")
BODY = 11.5
S = {
    "title": ParagraphStyle("title", fontName="TNR-B", fontSize=20, leading=26, alignment=TA_CENTER, spaceAfter=10),
    "subtitle": ParagraphStyle("subtitle", fontName="TNR", fontSize=13, leading=18, alignment=TA_CENTER,
                               textColor=GREY),
    "center": ParagraphStyle("center", fontName="TNR", fontSize=12, leading=17, alignment=TA_CENTER),
    "h1": ParagraphStyle("h1", fontName="TNR-B", fontSize=14.5, leading=19, spaceBefore=14, spaceAfter=6,
                         keepWithNext=1),
    "h2": ParagraphStyle("h2", fontName="TNR-B", fontSize=12.5, leading=17, spaceBefore=10, spaceAfter=4,
                         keepWithNext=1),
    "h3": ParagraphStyle("h3", fontName="TNR-BI", fontSize=11.5, leading=16, spaceBefore=6, spaceAfter=2,
                         keepWithNext=1),
    "body": ParagraphStyle("body", fontName="TNR", fontSize=BODY, leading=16, alignment=TA_JUSTIFY, spaceAfter=6),
    "q": ParagraphStyle("q", fontName="TNR-I", fontSize=BODY, leading=15.5, alignment=TA_JUSTIFY, leftIndent=18,
                        rightIndent=18, spaceAfter=7, textColor=GREY),
    "eq": ParagraphStyle("eq", fontName="TNR", fontSize=BODY, leading=17, alignment=TA_CENTER, spaceBefore=2,
                         spaceAfter=7),
    "item": ParagraphStyle("item", fontName="TNR", fontSize=BODY, leading=15.5, alignment=TA_JUSTIFY,
                           leftIndent=22, bulletIndent=4, spaceAfter=3),
    "cap": ParagraphStyle("cap", fontName="TNR", fontSize=10, leading=13, alignment=TA_CENTER, spaceBefore=3,
                          spaceAfter=10),
    "tcap": ParagraphStyle("tcap", fontName="TNR", fontSize=10, leading=13, alignment=TA_CENTER, spaceBefore=4,
                           spaceAfter=4),
    "cell": ParagraphStyle("cell", fontName="TNR", fontSize=9.5, leading=11.5, alignment=TA_LEFT),
    "cellh": ParagraphStyle("cellh", fontName="TNR-B", fontSize=9.5, leading=11.5, alignment=TA_LEFT),
    "ref": ParagraphStyle("ref", fontName="TNR", fontSize=10.5, leading=13.5, leftIndent=22, firstLineIndent=-22,
                          spaceAfter=3, alignment=TA_LEFT),
}
W = A4[0] - 2 * 2.3 * cm


def _jpeg(png: Path, max_w: int = 1700, quality: int = 80) -> Path:
    """Spectrogram PNGs are several MB each; both documents embed a JPEG copy."""
    from PIL import Image as PILImage
    out = REP / "_img" / (png.stem + ".jpg")
    out.parent.mkdir(parents=True, exist_ok=True)
    if not out.exists() or out.stat().st_mtime < png.stat().st_mtime:
        im = PILImage.open(png).convert("RGB")
        if im.width > max_w:
            im = im.resize((max_w, round(im.height * max_w / im.width)), PILImage.LANCZOS)
        im.save(out, "JPEG", quality=quality, optimize=True)
    return out


def _cell_text(v, fmt):
    if isinstance(v, (bool, np.bool_)):
        return "Yes" if v else "No"
    if isinstance(v, (int, np.integer)):
        return str(v).replace("-", "\u2212")
    if isinstance(v, float):
        s = f"{v:.1e}" if v != 0 and abs(v) < 1e-3 else fmt.format(v)
        return s.replace("-", "\u2212")
    return escape(str(v))


class Doc:
    """Collects the report once and keeps both a ReportLab story and a format-neutral block list."""

    def __init__(self, labels=None):
        self.story, self.blocks = [], []
        self.nfig = self.ntab = 0
        self.labels = labels or {}
        self.found = {}

    def ref(self, key):
        return str(self.labels.get(key, "?"))

    def _para(self, kind, text, style):
        self.blocks.append((kind, text))
        self.story.append(Paragraph(fix(text), S[style]))

    def h1(self, t):
        self._para("h1", t, "h1")

    def h2(self, t):
        self._para("h2", t, "h2")

    def h3(self, t):
        self._para("h3", t, "h3")

    def p(self, t):
        self._para("p", t, "body")

    def q(self, t):
        self._para("q", t, "q")

    def center(self, t, style="center"):
        self.blocks.append(("center", t, style))
        self.story.append(Paragraph(fix(t), S[style]))

    def eq(self, t, num=None):
        self.blocks.append(("eq", t, num))
        txt = fix(t) + (f"&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;({num})" if num else "")
        self.story.append(Paragraph(txt, S["eq"]))

    def items(self, its, numbered=False):
        self.blocks.append(("items", its, numbered))
        for i, it in enumerate(its, 1):
            self.story.append(Paragraph(fix(it), S["item"], bulletText=f"{i}." if numbered else "\u2022"))
        self.story.append(Spacer(1, 3))

    def space(self, h_cm):
        self.blocks.append(("space", h_cm))
        self.story.append(Spacer(1, h_cm * cm))

    def brk(self, need_cm=None):
        if need_cm is None:
            self.blocks.append(("pagebreak",))
            self.story.append(PageBreak())
        else:
            self.story.append(CondPageBreak(need_cm * cm))

    def fig(self, key, name, caption, width=1.0):
        path = _jpeg(FIG / name)
        iw, ih = ImageReader(str(path)).getSize()
        w = W * width
        h = w * ih / iw
        if h > 19 * cm:
            w, h = w * 19 * cm / h, 19 * cm
        self.nfig += 1
        self.found[key] = self.nfig
        cap = f"<b>Figure {self.nfig}:</b> {caption}"
        self.blocks.append(("fig", path, cap, w / cm))
        self.story.append(KeepTogether([Image(str(path), width=w, height=h), Paragraph(fix(cap), S["cap"])]))

    def table(self, key, df: pd.DataFrame, caption: str, col_widths=None, fmt="{:.3f}", index=True):
        d = df.reset_index() if index else df.copy()
        header = [escape(str(c)) for c in d.columns]
        body = [[_cell_text(v, fmt) for v in r] for r in d.astype(object).itertuples(index=False)]
        cw = col_widths or [W / len(header)] * len(header)
        self.ntab += 1
        self.found[key] = self.ntab
        cap = f"<b>Table {self.ntab}:</b> {caption}"
        self.blocks.append(("table", header, body, cap, [c / cm for c in cw]))
        rows = [[Paragraph(fix(h), S["cellh"]) for h in header]]
        rows += [[Paragraph(fix(c), S["cell"]) for c in r] for r in body]
        t = Table(rows, colWidths=cw, repeatRows=1, hAlign="CENTER")
        t.setStyle(TableStyle([("LINEABOVE", (0, 0), (-1, 0), 0.9, colors.black),
                               ("LINEBELOW", (0, 0), (-1, 0), 0.6, colors.black),
                               ("LINEBELOW", (0, -1), (-1, -1), 0.9, colors.black),
                               ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                               ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                               ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3)]))
        self.story.append(KeepTogether([Paragraph(fix(cap), S["tcap"]), t, Spacer(1, 9)]))

    def refs(self, items):
        self.blocks.append(("refs", items))
        for i, r in enumerate(items, 1):
            self.story.append(Paragraph(f"[{i}]&nbsp;&nbsp;{r}", S["ref"]))


# ----------------------------------------------------------------------------------------------- PDF
def _footer(canvas, doc):
    if doc.page == 1:
        return
    canvas.saveState()
    canvas.setFont("TNR", 9.5)
    canvas.setFillColor(GREY)
    canvas.drawString(2.3 * cm, 1.3 * cm, "Magnitude, Phase and Intelligibility of Speech")
    canvas.drawRightString(A4[0] - 2.3 * cm, 1.3 * cm, f"Page {doc.page}")
    canvas.setLineWidth(0.5)
    canvas.line(2.3 * cm, 1.7 * cm, A4[0] - 2.3 * cm, 1.7 * cm)
    canvas.restoreState()


def write_pdf(D, path, meta):
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=2.3 * cm, rightMargin=2.3 * cm, topMargin=2.2 * cm,
                            bottomMargin=2.4 * cm, title="Magnitude, Phase and Intelligibility of Speech",
                            author=meta.get("name", ""))
    doc.build(D.story, onFirstPage=_footer, onLaterPages=_footer)


# ----------------------------------------------------------------------------------------------- DOCX
_TOKEN = re.compile(r"(<[^>]+>|&nbsp;|&amp;|&lt;|&gt;|[^<&]+|&)")
_ENTITY = {"&nbsp;": "\u00a0", "&amp;": "&", "&lt;": "<", "&gt;": ">"}


def _runs(par, text, size=None, bold=False, italic=False):
    """Add text written with the small HTML-like markup of report_text.py as Word runs."""
    from docx.shared import Pt
    state = {"b": bold, "i": italic, "sub": False, "super": False}
    for tok in _TOKEN.findall(text):
        if tok.startswith("<"):
            tag = tok.strip("</>").split()[0].lower()
            if tag in state:
                state[tag] = not tok.startswith("</")
            elif tag == "br":
                par.add_run().add_break()
            continue
        r = par.add_run(_ENTITY.get(tok, tok))
        r.bold, r.italic = state["b"], state["i"]
        r.font.subscript, r.font.superscript = state["sub"], state["super"]
        r.font.name = "Times New Roman"
        if size:
            r.font.size = Pt(size)


def _times_style(style, size, bold=False, italic=False):
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor
    style.font.name, style.font.size = "Times New Roman", Pt(size)
    style.font.bold, style.font.italic = bold, italic
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(k), "Times New Roman")
    for k in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        fonts.attrib.pop(qn(k), None)


def write_docx(D, path):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.shared import Cm, Pt

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.3)
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.4)
    _times_style(doc.styles["Normal"], 12)
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.15
    for name, size in [("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 12)]:
        _times_style(doc.styles[name], size, bold=True, italic=name == "Heading 3")
        pf = doc.styles[name].paragraph_format
        pf.space_before, pf.space_after, pf.keep_with_next = Pt(12 if name == "Heading 1" else 8), Pt(4), True

    sizes = {"center": 12, "title": 20, "subtitle": 13}
    for b in D.blocks:
        kind = b[0]
        if kind in ("h1", "h2", "h3"):
            _runs(doc.add_heading(level=int(kind[1])), b[1])
        elif kind == "p":
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _runs(par, b[1])
        elif kind == "q":
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = par.paragraph_format.right_indent = Cm(0.6)
            _runs(par, b[1], italic=True)
        elif kind == "center":
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _runs(par, b[1], size=sizes.get(b[2], 12), bold=b[2] == "title")
        elif kind == "eq":
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _runs(par, b[1] + (f"\u00a0\u00a0\u00a0\u00a0({b[2]})" if b[2] else ""))
        elif kind == "items":
            for i, it in enumerate(b[1], 1):
                par = doc.add_paragraph()
                par.paragraph_format.left_indent = Cm(0.8)
                par.paragraph_format.first_line_indent = Cm(-0.5)
                _runs(par, (f"{i}.\u2002" if b[2] else "\u2022\u2002") + it)
        elif kind == "space":
            doc.add_paragraph().paragraph_format.space_after = Cm(b[1])
        elif kind == "pagebreak":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif kind == "fig":
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.keep_with_next = True
            par.add_run().add_picture(str(b[1]), width=Cm(b[3]))
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _runs(cap, b[2], size=10)
        elif kind == "table":
            header, body, cap_text, widths = b[1], b[2], b[3], b[4]
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.paragraph_format.keep_with_next = True
            _runs(cap, cap_text, size=10)
            t = doc.add_table(rows=1 + len(body), cols=len(header))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for i, row in enumerate([header] + body):
                for j, v in enumerate(row):
                    c = t.rows[i].cells[j]
                    c.width = Cm(widths[j])
                    c.paragraphs[0].paragraph_format.space_after = Pt(0)
                    _runs(c.paragraphs[0], v, size=9.5, bold=i == 0)
            doc.add_paragraph()
        elif kind == "refs":
            for i, r in enumerate(b[1], 1):
                par = doc.add_paragraph()
                par.paragraph_format.left_indent = Cm(0.8)
                par.paragraph_format.first_line_indent = Cm(-0.8)
                _runs(par, f"[{i}]\u2002{r}", size=11)
    doc.core_properties.title = "Magnitude, Phase and Intelligibility of Speech"
    doc.save(str(path))


# ----------------------------------------------------------------------------------------------- main
def _build(meta, labels):
    import report_text as RT
    D = Doc(labels)
    RT.build(D, meta)
    return D


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--name", default="Sarthak Tripathi")
    ap.add_argument("--usn", default="1MV23EC096")
    ap.add_argument("--department", default="Department of Electronics and Communication Engineering")
    ap.add_argument("--institution", default="Sir M. Visvesvaraya Institute of Technology, Bengaluru")
    ap.add_argument("--date", default="October 2026")
    ap.add_argument("--repo", default="https://github.com/Sarthak261-rgb/speech-phase-stoi")
    meta = vars(ap.parse_args())
    D = _build(meta, _build(meta, None).found)
    REP.mkdir(exist_ok=True)
    pdf, docx = REP / "Speech_Phase_STOI_Report.pdf", REP / "Speech_Phase_STOI_Report.docx"
    write_pdf(D, pdf, meta)
    write_docx(D, docx)
    print(f"wrote {pdf.name} ({pdf.stat().st_size / 1e6:.1f} MB) and {docx.name} "
          f"({docx.stat().st_size / 1e6:.1f} MB): {D.nfig} figures, {D.ntab} tables")


if __name__ == "__main__":
    main()
